"""Admin (HOD / Office Staff) editing: students, staff accounts, groups, Excel upload, requests, settings."""
import os, io, csv, json, uuid
from datetime import date
from flask import Blueprint, render_template, request, redirect, url_for, flash, g, abort, Response, send_file, current_app
from werkzeug.security import generate_password_hash as H, check_password_hash
import openpyxl
from models import *
from security import roles_required, audit, notify
import mentorship as ms

manage = Blueprint('manage', __name__)
ADM = ('hod', 'staff')
DEFAULT_PW = 'Welcome@123'

def clean(v, n=120): return (v or '').strip()[:n]

# ---------- students ----------
@manage.route('/admin/students')
@roles_required(*ADM)
def students():
    q = request.args.get('q', '').strip(); page = max(1, request.args.get('page', 1, type=int))
    qs = Student.query
    if q: qs = qs.filter(db.or_(Student.name.ilike(f'%{q}%'), Student.reg_no.ilike(f'%{q}%')))
    p = qs.order_by(Student.reg_no).paginate(page=page, per_page=25, error_out=False)
    cur = {s.id: ms.active_assignment(s.id) for s in p.items}
    return render_template('manage_students.html', p=p, q=q, cur=cur)

@manage.route('/admin/students/new', methods=['POST'])
@roles_required(*ADM)
def student_new():
    reg, name = clean(request.form['reg_no'], 20).upper(), clean(request.form['name'], 80)
    if not reg or not name: flash('Register number and name are required.', 'error')
    elif Student.query.filter_by(reg_no=reg).first() or User.query.filter_by(username=reg.lower()).first(): flash(f'{reg} already exists.', 'error')
    else:
        u = User(username=reg.lower(), name=name, role='student', pw_hash=H(DEFAULT_PW)); db.session.add(u); db.session.flush()
        db.session.add(Student(user_id=u.id, reg_no=reg, name=name, programme=clean(request.form.get('programme')) or 'M.Sc Statistics',
                               semester=clean(request.form.get('semester'), 5) or 'III', email=clean(request.form.get('email')), phone=clean(request.form.get('phone'), 20)))
        audit('student_create', f'{g.user.name} created {reg}'); db.session.commit()
        flash(f'Student {reg} created. Login: {reg.lower()} / {DEFAULT_PW}', 'ok')
    return redirect(url_for('manage.students'))

@manage.route('/admin/students/<reg>', methods=['GET', 'POST'])
@roles_required(*ADM)
def student_edit(reg):
    s = Student.query.filter_by(reg_no=reg).first_or_404()
    if request.method == 'POST':
        old = f'{s.name}|{s.programme}|{s.semester}'
        s.name = clean(request.form['name'], 80) or s.name; s.programme = clean(request.form['programme']); s.semester = clean(request.form['semester'], 5)
        s.email = clean(request.form['email']); s.phone = clean(request.form['phone'], 20); s.address = clean(request.form['address'], 200)
        if s.user_id: db.session.get(User, s.user_id).name = s.name
        audit('student_edit', f'{g.user.name} edited {reg}: {old} -> {s.name}|{s.programme}|{s.semester}'); db.session.commit()
        flash('Student details saved.', 'ok'); return redirect(url_for('manage.student_edit', reg=reg))
    recs = AcademicRecord.query.filter_by(student_id=s.id).order_by(AcademicRecord.period).all()
    return render_template('student_edit.html', s=s, recs=recs, a=ms.active_assignment(s.id))

@manage.route('/admin/students/<reg>/record', methods=['POST'])
@roles_required(*ADM)
def record_save(reg):
    s = Student.query.filter_by(reg_no=reg).first_or_404()
    try:
        per = int(request.form['period']); vals = [float(request.form[k]) for k in ('attendance', 'assessment', 'assignment_pct')]
        late = int(request.form.get('late') or 0)
        assert all(0 <= v <= 100 for v in vals) and per > 0 and late >= 0
    except (ValueError, AssertionError, KeyError): flash('Invalid values: percentages must be 0–100.', 'error'); return redirect(url_for('manage.student_edit', reg=reg))
    r = AcademicRecord.query.filter_by(student_id=s.id, period=per).first() or AcademicRecord(student_id=s.id, period=per)
    r.attendance, r.assessment, r.assignment_pct, r.late_submissions = *vals, late
    r.period_date = date.fromisoformat(request.form['period_date']) if request.form.get('period_date') else (r.period_date or date.today())
    db.session.add(r); audit('record_save', f'{g.user.name} saved period {per} for {reg}'); db.session.commit()
    flash(f'Period {per} saved.', 'ok'); return redirect(url_for('manage.student_edit', reg=reg))

@manage.route('/admin/students/<reg>/reset', methods=['POST'])
@roles_required(*ADM)
def student_reset(reg):
    s = Student.query.filter_by(reg_no=reg).first_or_404(); db.session.get(User, s.user_id).pw_hash = H(DEFAULT_PW)
    audit('password_reset', f'{g.user.name} reset password for {reg}'); db.session.commit()
    flash(f'Password reset to {DEFAULT_PW}. Ask the student to change it.', 'ok'); return redirect(url_for('manage.student_edit', reg=reg))

@manage.route('/admin/students/<reg>/delete', methods=['POST'])
@roles_required('hod')
def student_delete(reg):
    s = Student.query.filter_by(reg_no=reg).first_or_404()
    if request.form.get('confirm') != reg: flash('Type the register number to confirm deletion.', 'error'); return redirect(url_for('manage.student_edit', reg=reg))
    remove_student(s); audit('student_delete', f'{g.user.name} deleted {reg}'); db.session.commit()
    flash(f'{reg} and all linked records were deleted.', 'ok'); return redirect(url_for('manage.students'))

def remove_student(s):
    for M in (AcademicRecord, SupportRequest, Intervention, MentorAssignment): M.query.filter_by(student_id=s.id).delete()
    uid = s.user_id; db.session.delete(s)
    if uid: Notification.query.filter_by(user_id=uid).delete(); User.query.filter_by(id=uid).delete()

# ---------- faculty / staff accounts ----------
@manage.route('/admin/users', methods=['GET', 'POST'])
@roles_required(*ADM)
def users():
    allowed = ['mentor'] if g.user.role == 'staff' else ['mentor', 'staff', 'hod']
    if request.method == 'POST':
        un, name, role = clean(request.form['username'], 40).lower(), clean(request.form['name'], 80), request.form['role']
        if role not in allowed or not un or not name: flash('Invalid details or role not permitted.', 'error')
        elif User.query.filter_by(username=un).first(): flash('Username already taken.', 'error')
        else:
            db.session.add(User(username=un, name=name, role=role, email=clean(request.form.get('email')), pw_hash=H(request.form.get('password') or DEFAULT_PW)))
            audit('user_create', f'{g.user.name} created {role} {un}'); db.session.commit(); flash('Account created.', 'ok')
        return redirect(url_for('manage.users'))
    return render_template('manage_users.html', users=User.query.filter(User.role != 'student').order_by(User.role, User.name).all(), allowed=allowed)

@manage.route('/admin/users/<int:uid>', methods=['POST'])
@roles_required(*ADM)
def user_edit(uid):
    u = db.session.get(User, uid) or abort(404)
    if u.role == 'student' or (g.user.role == 'staff' and u.role != 'mentor'): abort(403)
    act = request.form.get('act')
    if act == 'save': u.name = clean(request.form['name'], 80) or u.name; u.email = clean(request.form.get('email'))
    elif act == 'reset': u.pw_hash = H(DEFAULT_PW); flash(f'Password reset to {DEFAULT_PW}.', 'ok')
    elif act == 'delete' and g.user.role == 'hod':
        if u.id == g.user.id: flash('You cannot delete your own account.', 'error')
        elif u.role == 'mentor' and (ms.students_of(u.id) or MentorGroup.query.filter_by(mentor_id=u.id, status='active').first()):
            flash('Reassign this mentor’s students and archive their groups first.', 'error')
        else: db.session.delete(u); flash('Account deleted.', 'ok')
    audit('user_' + str(act), f'{g.user.name} -> {u.username}'); db.session.commit(); return redirect(url_for('manage.users'))

# ---------- groups ----------
@manage.route('/admin/groups/<int:gid>', methods=['GET', 'POST'])
@roles_required(*ADM)
def group_edit(gid):
    gr = db.session.get(MentorGroup, gid) or abort(404)
    if request.method == 'POST':
        newm = int(request.form['mentor']); gr.name = clean(request.form['name'], 80) or gr.name
        gr.programme, gr.academic_year, gr.semester = clean(request.form['programme']), clean(request.form['year'], 10), clean(request.form['semester'], 5)
        gr.description = clean(request.form['description'], 200); gr.status = request.form['status'] if request.form['status'] in ('active', 'archived') else 'active'
        if newm != gr.mentor_id:  # mentor change: keep history for every member
            ids = [a.student_id for a in MentorAssignment.query.filter_by(group_id=gr.id, end_date=None)]
            gr.mentor_id = newm; ms.assign(ids, newm, gr.id, g.user.id, request.form.get('reason') or 'Group mentor changed'); notify(newm, f'You are now mentor of {gr.name}.')
        audit('group_edit', f'{g.user.name} edited {gr.name}'); db.session.commit(); flash('Group saved.', 'ok'); return redirect(url_for('admin.groups'))
    members = MentorAssignment.query.filter_by(group_id=gr.id, end_date=None).all()
    return render_template('group_edit.html', gr=gr, mentors=User.query.filter_by(role='mentor').all(), members=members)

@manage.route('/admin/assignment/<int:aid>/remove', methods=['POST'])
@roles_required(*ADM)
def unassign(aid):
    a = db.session.get(MentorAssignment, aid) or abort(404); a.end_date = date.today(); a.reason = (a.reason or '') + ' | ended: ' + clean(request.form.get('reason'), 80)
    audit('unassign', f'{g.user.name} ended assignment of {a.student.reg_no}'); db.session.commit(); return redirect(url_for('manage.group_edit', gid=a.group_id))

# ---------- Excel upload ----------
UP = {'students': (['reg_no', 'name'], ['programme', 'semester', 'email', 'phone']),
      'records': (['reg_no', 'period', 'attendance', 'assessment', 'assignment_pct'], ['late_submissions', 'period_date'])}

def upload_dir():
    d = os.path.join(current_app.instance_path, 'uploads'); os.makedirs(d, exist_ok=True); return d

@manage.route('/admin/upload', methods=['GET', 'POST'])
@roles_required(*ADM)
def upload():
    if request.method == 'GET': return render_template('upload.html')
    kind, f = request.form.get('kind'), request.files.get('file')
    if kind not in UP or not f or not f.filename.lower().endswith('.xlsx'): flash('Choose a data type and an .xlsx file.', 'error'); return redirect(url_for('manage.upload'))
    if g.user.role == 'staff' and kind == 'records': flash('Only the HOD can upload academic records.', 'error'); return redirect(url_for('manage.upload'))
    try: rows = list(openpyxl.load_workbook(f, read_only=True, data_only=True).active.iter_rows(values_only=True))
    except Exception: flash('That file could not be read as an Excel workbook.', 'error'); return redirect(url_for('manage.upload'))
    if len(rows) < 2: flash('The file has no data rows.', 'error'); return redirect(url_for('manage.upload'))
    head = [str(h or '').strip().lower() for h in rows[0]]; req, opt = UP[kind]
    miss = [c for c in req if c not in head]
    if miss: flash('Missing required column(s): ' + ', '.join(miss), 'error'); return redirect(url_for('manage.upload'))
    good, bad, seen = [], [], set(); regs = {s.reg_no: s.id for s in Student.query.all()}; new = upd = missing = 0
    for i, r in enumerate(rows[1:], start=2):
        d = {h: (str(v).strip() if v is not None else '') for h, v in zip(head, r)}
        if not any(d.values()): continue
        reg = d.get('reg_no', '').upper(); err = []
        if not reg: err.append('reg_no missing')
        if kind == 'students':
            if not d.get('name'): err.append('name missing')
            key = reg
            if reg in regs: upd += 1
            else: new += 1
        else:
            if reg and reg not in regs: err.append('unknown student (upload students first)')
            try:
                per = int(float(d['period'])); vals = [float(d[k]) for k in ('attendance', 'assessment', 'assignment_pct')]
                if not all(0 <= v <= 100 for v in vals) or per < 1: err.append('values out of range (percentages 0–100, period ≥ 1)')
            except (ValueError, KeyError): err.append('period/attendance/assessment/assignment_pct must be numbers')
            key = (reg, d.get('period'))
        if key in seen: err.append('duplicate row in file')
        seen.add(key); missing += sum(1 for c in req if not d.get(c))
        (bad if err else good).append({'row': i, 'data': d, 'errors': '; '.join(err)})
    tok = uuid.uuid4().hex
    json.dump({'kind': kind, 'user': g.user.id, 'good': [x['data'] for x in good], 'bad': bad}, open(os.path.join(upload_dir(), tok + '.json'), 'w'))
    return render_template('upload_preview.html', tok=tok, kind=kind, received=len(good) + len(bad), good=good, bad=bad, new=new, upd=upd, missing=missing)

def load_job(tok):
    if not tok.isalnum(): abort(404)
    p = os.path.join(upload_dir(), tok + '.json')
    if not os.path.exists(p): abort(404)
    j = json.load(open(p))
    if j['user'] != g.user.id: abort(403)
    return j, p

@manage.route('/admin/upload/<tok>/errors.csv')
@roles_required(*ADM)
def upload_errors(tok):
    j, _ = load_job(tok); out = io.StringIO(); w = csv.writer(out); w.writerow(['Excel row', 'Errors', 'Data'])
    for b in j['bad']: w.writerow([b['row'], b['errors'], json.dumps(b['data'])])
    return Response(out.getvalue(), mimetype='text/csv', headers={'Content-Disposition': 'attachment; filename=upload_errors.csv'})

@manage.route('/admin/upload/<tok>/confirm', methods=['POST'])
@roles_required(*ADM)
def upload_confirm(tok):
    j, p = load_job(tok); n = u = 0
    for d in j['good']:
        reg = d['reg_no'].upper()
        if j['kind'] == 'students':
            s = Student.query.filter_by(reg_no=reg).first()
            if s: s.name, s.programme, s.semester = d['name'], d.get('programme') or s.programme, d.get('semester') or s.semester; u += 1
            elif User.query.filter_by(username=reg.lower()).first(): continue
            else:
                us = User(username=reg.lower(), name=d['name'], role='student', pw_hash=H(DEFAULT_PW)); db.session.add(us); db.session.flush()
                db.session.add(Student(user_id=us.id, reg_no=reg, name=d['name'], programme=d.get('programme') or 'M.Sc Statistics', semester=d.get('semester') or 'III',
                                       email=d.get('email', ''), phone=d.get('phone', ''))); n += 1
            db.session.flush()
        else:
            s = Student.query.filter_by(reg_no=reg).first(); per = int(float(d['period']))
            r = AcademicRecord.query.filter_by(student_id=s.id, period=per).first()
            if r: u += 1
            else: r = AcademicRecord(student_id=s.id, period=per); n += 1
            r.attendance, r.assessment, r.assignment_pct = float(d['attendance']), float(d['assessment']), float(d['assignment_pct'])
            r.late_submissions = int(float(d.get('late_submissions') or 0))
            try: r.period_date = date.fromisoformat(d['period_date'][:10])
            except (ValueError, KeyError): r.period_date = r.period_date or date.today()
            db.session.add(r)
    audit('upload', f'{g.user.name} imported {j["kind"]}: {n} new, {u} updated, {len(j["bad"])} rejected'); db.session.commit(); os.remove(p)
    flash(f'Import complete: {n} new, {u} updated, {len(j["bad"])} rows rejected (see error report if needed).', 'ok'); return redirect(url_for('manage.upload'))

@manage.route('/admin/upload/template/<kind>')
@roles_required(*ADM)
def template(kind):
    if kind not in UP: abort(404)
    wb = openpyxl.Workbook(); ws = wb.active; ws.append(UP[kind][0] + UP[kind][1])
    ws.append(['P252701', 'Arun K', 'M.Sc Statistics', 'III', 'arun@example.com', '9999999999'] if kind == 'students' else ['P252701', 1, 82, 70, 90, 0, '2026-09-01'])
    b = io.BytesIO(); wb.save(b); b.seek(0); return send_file(b, as_attachment=True, download_name=f'{kind}_template.xlsx')

# ---------- support requests (admin view) ----------
@manage.route('/admin/requests', methods=['GET', 'POST'])
@roles_required(*ADM)
def requests_admin():
    if request.method == 'POST':
        r = db.session.get(SupportRequest, int(request.form['rid'])) or abort(404)
        if request.form['status'] in ('Submitted', 'Assigned', 'In Progress', 'Follow-up', 'Resolved'):
            r.status = request.form['status']; notify(r.student.user_id, f'Your support request is now: {r.status}.'); audit('request_status', f'{r.student.reg_no} -> {r.status}'); db.session.commit()
        return redirect(url_for('manage.requests_admin'))
    return render_template('requests_admin.html', reqs=SupportRequest.query.order_by(SupportRequest.created.desc()).limit(200).all(),
                           assign={r.student_id: ms.active_assignment(r.student_id) for r in SupportRequest.query.all()})

# ---------- settings ----------
@manage.route('/admin/settings', methods=['GET', 'POST'])
@roles_required('hod')
def settings():
    demo_s, demo_u = Student.query.filter_by(is_demo=True).count(), User.query.filter_by(is_demo=True).count()
    if request.method == 'POST' and request.form.get('confirm') == 'REMOVE DEMO':
        for s in Student.query.filter_by(is_demo=True).all(): remove_student(s)
        for u in User.query.filter_by(is_demo=True).all():
            if u.id != g.user.id and u.role != 'student':
                MentorGroup.query.filter_by(mentor_id=u.id).delete(); db.session.flush(); Notification.query.filter_by(user_id=u.id).delete(); db.session.delete(u)
        audit('demo_removed', g.user.name); db.session.commit(); flash('Demo data removed. Your own records were not touched.', 'ok'); return redirect(url_for('manage.settings'))
    return render_template('settings.html', demo_s=demo_s, demo_u=demo_u)
